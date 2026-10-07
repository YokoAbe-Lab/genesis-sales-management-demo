"""Export a frozen approved Run; never reread or reanalyze its source."""
import base64
import html
import io
import json
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import analysis_core as core

ROOT=Path(__file__).resolve().parent
NODE=Path('node')

def set_excel_print_layout(path,analysis):
    """Set print-only OOXML metadata after artifact-tool writes the new workbook."""
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    tag=lambda name:'{'+ns+'}'+name
    ET.register_namespace('',ns)
    with zipfile.ZipFile(path) as src:
        files={i.filename:src.read(i.filename) for i in src.infolist()}
    for number in (1,2):
        name=f'xl/worksheets/sheet{number}.xml';root=ET.fromstring(files[name])
        prop=root.find(tag('sheetPr'))
        if prop is None:prop=ET.Element(tag('sheetPr'));root.insert(0,prop)
        setting=prop.find(tag('pageSetUpPr'))
        if setting is None:setting=ET.SubElement(prop,tag('pageSetUpPr'))
        setting.set('fitToPage','1')
        for elem in ['pageMargins','pageSetup','headerFooter']:
            old=root.find(tag(elem))
            if old is not None:root.remove(old)
        # These elements belong after printOptions and before breaks/drawings.
        later={'rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
        pos=next((i for i,x in enumerate(root) if x.tag.rsplit('}',1)[-1] in later),len(root))
        margins=ET.Element(tag('pageMargins'),left='0.25',right='0.25',top='0.4',bottom='0.4',header='0.15',footer='0.15')
        setup=ET.Element(tag('pageSetup'),paperSize='9',orientation='landscape',fitToWidth='1',fitToHeight='0')
        footer=ET.Element(tag('headerFooter'));ET.SubElement(footer,tag('oddFooter')).text='&LGenesis 分析資料&R&P / &N'
        for offset,element in enumerate((margins,setup,footer)):root.insert(pos+offset,element)
        files[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    root=ET.fromstring(files['xl/workbook.xml']);names=root.find(tag('definedNames'))
    if names is None:
        names=ET.Element(tag('definedNames'));sheets=root.find(tag('sheets'));root.insert(list(root).index(sheets)+1,names)
    for n in list(names):
        if n.get('name') in ['_xlnm.Print_Area','_xlnm.Print_Titles']:names.remove(n)
    end=8+len(analysis['series'])+len(analysis['sections'])+3+1
    for index,(sheet,area) in enumerate([('集計結果',f'$A$1:$D${end}'),('元データと条件',f'$A$1:$J${7+len(analysis["evidence"])}')]):
        ET.SubElement(names,tag('definedName'),name='_xlnm.Print_Area',localSheetId=str(index)).text=f"'{sheet}'!{area}"
    ET.SubElement(names,tag('definedName'),name='_xlnm.Print_Titles',localSheetId='1').text="'元データと条件'!$7:$7"
    files['xl/workbook.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
        for name,data in files.items():dst.writestr(name,data)
def lines(r,approval,artifact):
    a=r['analysis']
    return [('案件',r['projectName']),('RequestID',r['requestId']),('RunID',r['runId']),('ApprovalID',approval['id']),('ArtifactID',artifact),('AgentRunID',r['agentRunId']+' LocalRule AI通信なし'),
            ('QuestionID',', '.join(q['questionId'] for q in r['questionAnswers'])),('TestCaseID',', '.join(r['testCaseIds'])),
            ('計算日時',a['calculatedAt']),('元ファイル',a['source']['path']),('シートまたは表',a['source']['table']),('SHA256',a['source']['sha256']),('条件',a['conditions']),('データ点検',f"{a['qualityRows']}行に注意。集計対象 {a['summary']['count']}行。"),('実行と試験の範囲','架空サンプル。実データ・Access・AI接続の試験を示しません。')]
def export_files(r,approval,out):
    a=r['analysis'];picture=core.chart_png(a);ext=out.suffix[1:];meta=lines(r,approval,out.stem)
    if ext=='png':out.write_bytes(picture);return
    if ext=='json':
        out.write_text(json.dumps({'artifactId':out.stem,'approval':approval,'run':r},ensure_ascii=False,indent=2),encoding='utf-8');return
    if ext=='html':
        e=html.escape
        articles=''.join(f'<section><p class="kind">{kind}</p><h2>{title}</h2><p>{e(a["finalText"][key]).replace(chr(10),"<br>")}</p></section>' for key,kind,title in core.SECTIONS)
        table=''.join(f'<tr><th>{e(k)}</th><td>{e(v)}</td></tr>' for k,v in meta)
        body=f'''<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{e(a['source']['name'])}の分析報告書</title>
        <style>body{{font:18px/1.9 Meiryo,sans-serif;color:#153042;max-width:1100px;margin:40px auto;padding:25px}}h1,h2{{color:#122b3c}}h1{{font-size:30px}}h2{{font-size:23px;margin:0 0 12px}}p{{overflow-wrap:anywhere}}img{{width:100%;height:auto}}.kind{{color:#9a5c23;margin:30px 0 4px}}table{{border-collapse:collapse;width:100%;font-size:14px;table-layout:fixed}}th,td{{border:1px solid #a9bfcd;padding:10px;text-align:left;overflow-wrap:anywhere}}th{{width:22%;background:#ecf4f9}}section{{break-inside:avoid}}html{{background:#f4f7f9}}.kind{{display:inline-block;padding:2px 8px;border-radius:4px;background:#fff3e6;color:#784719}}@media print{{@page{{size:A4 landscape;margin:12mm}}body{{margin:0;padding:0;max-width:none;font-size:10.5pt;line-height:1.55;print-color-adjust:exact;-webkit-print-color-adjust:exact}}img{{display:block;max-height:92mm;object-fit:contain}}.kind{{margin:12px 0 4px}}h2{{font-size:15pt}}}}</style>
        <h1>{e(a['source']['name'])}の分析報告書</h1><p>架空データによる確認用資料です。図表と説明を人が確認した時点の結果を保存しています。</p><p>{e(r['runId'])} ／ {e(approval['id'])}</p><img alt="選択されたグラフ" src="data:image/png;base64,{base64.b64encode(picture).decode()}">
        {articles}<h2>条件と根拠の一覧</h2><table>{table}</table><p>文章はローカルの定型処理で生成し、利用者が確認・修正した内容です。因果関係の証明ではありません。</p></html>'''
        out.write_text(body,encoding='utf-8');return
    if ext=='docx':
        from docx import Document
        from docx.shared import Cm,Pt,RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
        for name in ['Normal','Title','Heading 1','Heading 2']:
            st=doc.styles[name];st.font.name='Meiryo';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Meiryo')
        for st in doc.styles:
            for border in st._element.xpath('.//w:pBdr'):
                border.getparent().remove(border)
        doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.15
        doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Title'].font.size=Pt(23);doc.styles['Heading 1'].font.size=Pt(15)
        doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
        doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
        doc.add_paragraph(a['source']['name']+'の分析報告書','Title')
        doc.add_paragraph('架空データによる確認用資料です。図表と文章を確認した時点の結果をまとめます。実務上の原因や対策の有効性を証明する資料ではありません。')
        doc.add_paragraph(r['runId']+'　承認 '+approval['id'])
        doc.add_picture(io.BytesIO(picture),width=Cm(17))
        doc.add_heading('事実 現在の状況',1);doc.add_paragraph(a['finalText']['situation'])
        doc.add_heading('事実 最大の特徴',1);doc.add_paragraph(a['finalText']['feature'])
        doc.add_page_break()
        for key,kind,title in core.SECTIONS:
            if key in ['situation','feature','source']:continue
            doc.add_heading(kind+' '+title,1);doc.add_paragraph(a['finalText'][key])
        doc.add_page_break()
        doc.add_heading('使用したファイルと条件',1);doc.add_paragraph(a['finalText']['source'])
        doc.add_heading('根拠と追跡情報',1)
        tab=doc.add_table(rows=0,cols=2);tab.autofit=False;tab.columns[0].width=Cm(3);tab.columns[1].width=Cm(14)
        for k,v in meta:
            cells=tab.add_row().cells;cells[0].text=k;cells[1].text=v
            for idx,cell in enumerate(cells):
                pr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                if idx==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8F1F7');pr.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1
                    for run in p.runs:run.font.size=Pt(9)
        for p in list(doc.paragraphs)+[p for row in tab.rows for cell in row.cells for p in cell.paragraphs]:
            p.paragraph_format.keep_together=True
            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
        footer=sec.footer.paragraphs[0];footer.add_run('Genesis 分析資料　')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
        doc.core_properties.author='AbeLab';doc.core_properties.title=a['source']['name']+'の分析報告書'
        doc.save(out);return
    if ext=='pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Image,Table,TableStyle
        from reportlab.lib.pagesizes import A4,landscape
        pdfmetrics.registerFont(TTFont('ReportJP',str(core.FONT),subfontIndex=0))
        page_width,page_height=landscape(A4)
        margin=36
        content_width=page_width-2*margin-12
        body=ParagraphStyle('body',fontName='ReportJP',fontSize=10.5,leading=16,wordWrap='CJK',spaceAfter=7,textColor=colors.HexColor('#153244'))
        title=ParagraphStyle('title',parent=body,fontSize=20,leading=27,spaceAfter=10,textColor=colors.HexColor('#122b3c'))
        heading=ParagraphStyle('head',parent=body,fontSize=13,leading=19,spaceBefore=9,spaceAfter=5,keepWithNext=True,textColor=colors.HexColor('#122b3c'),backColor=colors.HexColor('#edf4f8'),borderPadding=4)
        important=ParagraphStyle('important',parent=heading,backColor=colors.HexColor('#fff3e6'))
        small=ParagraphStyle('small',parent=body,fontSize=8.5,leading=13)
        confirmed=ParagraphStyle('confirmed',parent=small,backColor=colors.HexColor('#e9f5ee'),borderPadding=3)
        esc=lambda t:html.escape(str(t)).replace('\n','<br/>')
        story=[Paragraph(esc(a['source']['name']+'の分析報告書'),title),Paragraph('架空データによる確認用資料です。人が確認した時点の図表と説明を保存しています。',body),Paragraph(esc(r['runId']+' ／ '+approval['id']),confirmed),Image(io.BytesIO(picture),width=560,height=560*520/1100)]
        for key in ['situation','feature']:
            story.extend([Paragraph('事実 '+('現在の状況' if key=='situation' else '最大の特徴'),important if key=='feature' else heading),Paragraph(esc(a['finalText'][key]),body)])
        for key,kind,name in core.SECTIONS:
            if key in ['situation','feature','source']:continue
            story.extend([Paragraph(esc(kind+' '+name),important if key in ['caution','cause','limits'] else heading),Paragraph(esc(a['finalText'][key]),body)])
        story.extend([Paragraph('使用したファイルと条件',heading),Paragraph(esc(a['finalText']['source']),small),Paragraph('根拠と追跡情報',heading)])
        table=Table([[Paragraph(esc(k),small),Paragraph(esc(v),small)] for k,v in meta],colWidths=[100,content_width-100])
        table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#a9bfcd')),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#E8F1F7')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        story.append(table)
        def footer(canvas,doc):
            canvas.saveState()
            canvas.setFillColor(colors.HexColor('#f4f7f9'));canvas.rect(0,0,page_width,page_height,stroke=0,fill=1)
            canvas.setFillColor(colors.HexColor('#153244'));canvas.setFont('ReportJP',8)
            canvas.drawString(margin,19,'Genesis 分析資料');canvas.drawRightString(page_width-margin,19,str(doc.page))
            canvas.restoreState()
        SimpleDocTemplate(str(out),pagesize=(page_width,page_height),leftMargin=margin,rightMargin=margin,topMargin=28,bottomMargin=32,title=a['source']['name']+'の分析報告書',author='AbeLab').build(story,onFirstPage=footer,onLaterPages=footer);return
    if ext=='xlsx':
        payload={'run':r,'approval':approval,'metadata':meta}
        intermediate=out.with_suffix('.input.json');intermediate.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
        result=subprocess.run([str(NODE),str(ROOT/'export_excel.mjs'),str(intermediate),str(out)],capture_output=True,timeout=120,creationflags=subprocess.CREATE_NO_WINDOW)
        if result.returncode: raise RuntimeError('EXCEL_EXPORT_FAILED')
        set_excel_print_layout(out,a)
        return
    raise ValueError('Unsupported format')
